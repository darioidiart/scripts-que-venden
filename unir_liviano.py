import glob
from openpyxl import Workbook, load_workbook

archivos = glob.glob("*.xlsx")
wb_final = Workbook()
ws_final = wb_final.active
primera_vez = True

for archivo in archivos:
    if archivo == "RESULTADO_FINAL.xlsx":
        continue
    wb = load_workbook(archivo)
    ws = wb.active
    for i, fila in enumerate(ws.iter_rows(values_only=True)):
        if i==0 and not primera_vez:
            continue # saltea encabezado
        ws_final.append(fila)
    primera_vez = False
    print(f"+ {archivo}")

wb_final.save("RESULTADO_FINAL.xlsx")
print("LISTO! RESULTADO_FINAL.xlsx creado")
